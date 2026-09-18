"""Turns the Belgian customs systems' rule and code-list workbooks into JSON (see sync-be-docs.ps1).

  IDMS  IDMS_VRE.xlsx -> idms-rules.json, IDMS_codelist.xlsx -> idms-codelists.json
  AES   AES_rules.xlsx -> aes-rules.json, AES_national_rules.xlsx -> aes-national-rules.json
  NCTS  NCTS_P5_rules.xlsx -> ncts-p5-rules.json, NCTS_CLBE*.xlsx -> ncts-national-codelists.json

WHY THE SHEET MATTERS MORE THAN THE RULE TEXT: IDMS and AES keep switched-off rules on a sheet of
their own. BE0211 and BE0212 read as hard requirements, and both sit on IDMS's "RulesInactive"
sheet in the 2026-09-08 workbook. A consumer that loads the rules without their sheet cannot tell
the two apart, so every rule carries the sheet it came from, and a "status" where the sheet's name
states one. Where it does not (NCTS publishes no inactive sheet) there is no status: guessing
"active" would be exactly the mistake this exists to prevent.

Usage: python parse-be-docs.py <IDMS|AES|NCTS> <folder>
"""

import json
import re
import sys
from pathlib import Path

import openpyxl


def cell(v):
    """A cell as text, or None. Codes arrive as numbers where Excel guessed (CL093's 22, 43…)."""
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip()
    return s or None


def rows(ws):
    return [[cell(c) for c in r] for r in ws.iter_rows(values_only=True)]


def workbook(path):
    return openpyxl.load_workbook(path, read_only=True, data_only=True)


def write(folder, name, data):
    (folder / name).write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")


# ─── Rules ────────────────────────────────────────────────────────────────────

FIELD = {
    "id": "id", "r/c code": "code", "precondition": "precondition",
    "functional description": "functionalDescription", "technical description": "technicalDescription",
    "error code": "errorCode", "error reason": "errorReason", "description": "description",
    "error fr": "errorFr", "error nl": "errorNl", "error de": "errorDe", "error en": "errorEn",
    "domain": "domain", "notes": "notes", "implemenation by ibm": "implementationByIbm",
    "implemented as drools": "implementedAsDrools", "ibm/vre": "ibmVre",
}

CODE_HEADERS = ("r/c code", "error reason")


def field(header):
    h = header.strip().lower()
    return FIELD.get(h) or re.sub(r"[^A-Za-z0-9]+(.)", lambda m: m.group(1).upper(), h)


def header_row(body):
    """Index of the header: the first of the top rows naming the rule-code column."""
    for i, r in enumerate(body[:5]):
        if any(c and c.lower() in CODE_HEADERS for c in r):
            return i
    return None


def read_rules(ws, status=None, fallback_header=None):
    body = rows(ws)
    at = header_row(body)
    if at is None:
        if fallback_header is None:
            raise SystemExit(f"Sheet '{ws.title}' has no header and none to borrow")
        header, data = fallback_header, body
    else:
        header, data = body[at], body[at + 1:]

    lowered = [(h or "").lower() for h in header]
    code_col = next(lowered.index(h) for h in CODE_HEADERS if h in lowered)

    out = []
    for r in data:
        code = r[code_col] if code_col < len(r) else None
        # A row with no rule code is a note, a draft or a spacer, not a rule.
        if not code or code.lower() in CODE_HEADERS:
            continue
        rule = {"sheet": ws.title}
        if status:
            rule["status"] = status
        for i, h in enumerate(header):
            if h and i < len(r) and r[i] is not None:
                rule[field(h)] = r[i]
        rule["code"] = code
        out.append(rule)
    return out, header


def read_changes(ws, skip):
    """A change log: [rule, change] rows after `skip` heading rows."""
    return [{"code": r[0], "change": r[1] if len(r) > 1 else None}
            for r in rows(ws)[skip:] if r and r[0]]


def rules_file(path, sheets, changes=None, minimum=100):
    """sheets: [(sheet name, status or None, borrow the header of this sheet or None)]."""
    wb = workbook(path)
    rules, headers = [], {}
    for name, status, borrow in sheets:
        if name not in wb.sheetnames:
            raise SystemExit(f"{path.name} has no '{name}' sheet: {wb.sheetnames}")
        found, header = read_rules(wb[name], status, headers.get(borrow))
        headers[name] = header
        rules += found
    if len(rules) < minimum:
        raise SystemExit(f"{path.name} yielded only {len(rules)} rules -- not the workbook we expect")

    result = {"counts": {}, "rules": rules}
    for r in rules:
        k = r.get("status") or r["sheet"]
        result["counts"][k] = result["counts"].get(k, 0) + 1
    if changes and changes[0] in wb.sheetnames:
        result["changes"] = read_changes(wb[changes[0]], changes[1])
    return result


# ─── Code lists ───────────────────────────────────────────────────────────────

DESCRIPTION_LANG = {"description en": "en", "description fr": "fr", "description nl": "nl",
                    "descriptionnl": "nl", "description de": "de", "description": "en",
                    "en": "en", "fr": "fr", "nl": "nl", "de": "de"}

# List-level columns of a code list exported from JSON and flattened into a sheet (NCTS CLBE213):
# they describe the list, sit on its first row only, and are not attributes of an entry.
LIST_LEVEL = {"name", "codelists", "lastupdate", "code", "entries", "languages", "codes"}


def value_column(r):
    return next((i for i, c in enumerate(r) if c and c.lower() == "value"), None)


def entries_below(body, head_at):
    """Entries under a header row that names a "Value" column; one-cell rows are section headings."""
    header = body[head_at]
    value_col = value_column(header)
    entries, section = [], None
    for r in body[head_at + 1:]:
        filled = [c for c in r if c]
        if not filled:
            continue
        value = r[value_col] if value_col < len(r) else None
        if not value:
            if len(filled) == 1:
                section = filled[0]
            continue
        entry = {"value": value}
        for i, h in enumerate(header):
            if not h or i == value_col or i >= len(r) or r[i] is None:
                continue
            lang = DESCRIPTION_LANG.get(h.lower())
            if lang:
                entry.setdefault("descriptions", {})[lang] = r[i]
            elif h.lower() in LIST_LEVEL:
                continue            # the list's own name/code, repeated or on the first row only
            elif h == "CCI filter":
                entry["cciFilter"] = r[i]
            elif h.startswith(("Appendice", "Appendix")):
                entry["appendix"] = r[i]
            else:
                entry[h] = r[i]
        if section:
            entry["section"] = section
        entries.append(entry)
    return entries


def idms_codelists(path):
    wb = workbook(path)
    index = {}
    for r in rows(wb["Code list"])[1:]:
        if r and r[0]:
            index[r[0].strip()] = {"name": r[1] if len(r) > 1 else None,
                                   "cciFilter": r[2] if len(r) > 2 else None,
                                   "nationalValue": r[3] if len(r) > 3 else None}
    lists = {}
    for ws in wb.worksheets:
        if ws.title == "Code list":
            continue
        body = rows(ws)
        code = (body[0][0] if body and body[0] and body[0][0] else ws.title).strip()
        head_at = next((i for i, r in enumerate(body) if value_column(r) is not None), None)
        if head_at is None:
            continue
        # Above the header: the list code, "Home", and the data elements it serves.
        data_elements = [c for r in body[:head_at] for c in r[1:] if c and c != "Home" and c[:1].isdigit()]
        lists[code] = {**index.get(code, {}), "dataElements": data_elements,
                       "entries": entries_below(body, head_at)}
    if len(lists) < 20:
        raise SystemExit(f"{path.name} yielded only {len(lists)} lists -- not the workbook we expect")
    return {"lists": lists}


def national_codelists(paths):
    """One small workbook per list (NCTS_CLBE009.xlsx …), a Value column and a description."""
    lists = {}
    for p in paths:
        code = re.search(r"CLBE\d{3}", p.name).group(0)
        body = rows(workbook(p).worksheets[0])
        head_at = next((i for i, r in enumerate(body) if value_column(r) is not None), None)
        if head_at is None:
            raise SystemExit(f"{p.name} has no 'Value' column")
        lists[code] = {"entries": entries_below(body, head_at)}
    if not lists:
        raise SystemExit("No national code-list workbooks found")
    return {"lists": lists}


# ─── Per system ───────────────────────────────────────────────────────────────

def idms(folder):
    rules = rules_file(folder / "IDMS_VRE.xlsx",
                       [("Feuil1", "active", None),
                        ("RulesInactive", "inactive", "Feuil1"),
                        ("TemporaryBlockingRules", "temporary-blocking", "Feuil1")],
                       changes=("Changes", 1))
    write(folder, "idms-rules.json", rules)
    lists = idms_codelists(folder / "IDMS_codelist.xlsx")
    write(folder, "idms-codelists.json", lists)
    return [("idms-rules.json", rules["counts"]),
            ("idms-codelists.json", {"lists": len(lists["lists"])})]


def aes(folder):
    sheets = [("active VRE rules", "active", None),
              ("inactive rules", "inactive", "active VRE rules"),
              ("Other active rules", "active", None),
              ("Missing documents filter", None, None)]
    out = []
    for source, target in (("AES_rules.xlsx", "aes-rules.json"),
                           ("AES_national_rules.xlsx", "aes-national-rules.json")):
        rules = rules_file(folder / source, sheets, changes=("Release notes", 2))
        write(folder, target, rules)
        out.append((target, rules["counts"]))
    return out


def ncts(folder):
    rules = rules_file(folder / "NCTS_P5_rules.xlsx",
                       [("Rules and Conditions (for VRE)", None, None),
                        ("Rules and Conditions (2)", None, None)])
    write(folder, "ncts-p5-rules.json", rules)
    lists = national_codelists(sorted(folder.glob("NCTS_CLBE*.xlsx")))
    write(folder, "ncts-national-codelists.json", lists)
    return [("ncts-p5-rules.json", rules["counts"]),
            ("ncts-national-codelists.json", {"lists": len(lists["lists"])})]


def main():
    system, folder = sys.argv[1].upper(), Path(sys.argv[2])
    for name, counts in {"IDMS": idms, "AES": aes, "NCTS": ncts}[system](folder):
        print(f"{name}: " + ", ".join(f"{v} {k}" for k, v in counts.items()))


if __name__ == "__main__":
    main()
